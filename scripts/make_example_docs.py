from pathlib import Path
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table

out = Path("examples")
out.mkdir(exist_ok=True)

title = "Quarterly Data Quality Report (demo)"
paras = [
    "This is a synthetic document generated for testing Nydra. "
    "It contains no real data.",
    "Summary: the order dataset had 3 percent missing values and a few "
    "duplicate rows. Cleaning removed the duplicates and filled gaps with "
    "median values.",
    "Contact for questions: analyst@example.com, phone 555-0142.",
    "Next steps: review outliers in the price column, then retrain the "
    "model and compare results with the previous quarter.",
]
table = [["Metric", "Before", "After"],
         ["Missing values", "3.1%", "0%"],
         ["Duplicate rows", "5", "0"],
         ["Outliers flagged", "4", "4"]]

# PDF
styles = getSampleStyleSheet()
story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
for p in paras:
    story += [Paragraph(p, styles["BodyText"]), Spacer(1, 8)]
story.append(Table(table))
SimpleDocTemplate(str(out / "demo_report.pdf"), pagesize=A4,
                  title=title, author="Nydra demo").build(story)

# DOCX
doc = Document()
doc.core_properties.author = "Nydra demo"
doc.core_properties.title = title
doc.core_properties.comments = ""
doc.add_heading(title, 0)
for p in paras:
    doc.add_paragraph(p)
t = doc.add_table(rows=len(table), cols=3)
for r, row in enumerate(table):
    for c, val in enumerate(row):
        t.cell(r, c).text = val
doc.save(str(out / "demo_report.docx"))
print("OK")
